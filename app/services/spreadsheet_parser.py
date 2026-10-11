"""
Spreadsheet Parser Service
Reads, normalizes, and validates Excel (.xlsx) and CSV files for bulk certificate jobs.
"""

import csv
import io
import logging
import zipfile
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils.exceptions import InvalidFileException
from pydantic import ValidationError

from app.core.config import settings
from app.core.errors import (
    BatchSizeExceededError,
    InvalidFileFormatError,
    SpreadsheetParsingError,
)
from app.schemas.recipient import RecipientInput

logger = logging.getLogger("bulk_certificate_generator.spreadsheet")

# Column header normalization synonyms
NAME_VARIANTS = {
    "recipient_name",
    "recipientname",
    "name",
    "fullname",
    "full_name",
    "participant_name",
    "participantname",
    "participant",
    "student_name",
    "studentname",
    "student",
    "attendee_name",
    "attendeename",
    "attendee",
    "candidate_name",
    "candidatename",
    "candidate",
}

EMAIL_VARIANTS = {
    "recipient_email",
    "recipientemail",
    "email",
    "email_address",
    "emailaddress",
    "e_mail",
    "emailaddr",
    "contact_email",
    "contactemail",
    "mail",
}

REMARKS_VARIANTS = {
    "custom_remarks",
    "customremarks",
    "remarks",
    "remark",
    "notes",
    "note",
    "grade",
    "honors",
    "honor",
    "comments",
    "comment",
    "custom_notes",
    "custom_note",
    "award_title",
    "award",
}


def _clean_header(val: Any) -> str:
    """Normalizes header string for fuzzy matching."""
    if val is None:
        return ""
    return str(val).strip().lower().replace(" ", "_").replace("-", "").replace("_", "")


class SpreadsheetParser:
    """Parser and validator for recipient spreadsheets."""

    def parse_recipients_file(self, file_bytes: bytes, filename: str) -> list[RecipientInput]:
        """
        Parses an Excel (.xlsx) or CSV file from in-memory bytes and returns a validated
        list of RecipientInput objects.
        """
        if not file_bytes or len(file_bytes) == 0:
            raise InvalidFileFormatError("The uploaded file is empty (0 bytes).")

        ext = Path(filename).suffix.lower()
        if ext == ".xls":
            raise InvalidFileFormatError(
                "Legacy Excel format (.xls) is not supported. Please save and upload as modern Excel (.xlsx) or CSV (.csv)."
            )
        if ext not in (".xlsx", ".csv"):
            raise InvalidFileFormatError(
                f"Unsupported file format '{ext or 'unknown'}'. Please upload an Excel (.xlsx) or CSV (.csv) file."
            )

        if ext == ".xlsx":
            rows = self._read_xlsx_rows(file_bytes)
        else:
            rows = self._read_csv_rows(file_bytes)

        return self._extract_and_validate_recipients(rows)

    def _read_xlsx_rows(self, file_bytes: bytes) -> list[list[Any]]:
        """Reads rows from an Excel .xlsx workbook."""
        try:
            wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=True)
        except (InvalidFileException, zipfile.BadZipFile, Exception) as exc:
            logger.warning("Failed to open Excel workbook: %s", exc)
            raise InvalidFileFormatError(
                "Unable to read Excel file: the file appears corrupted or is not a valid .xlsx workbook."
            ) from exc

        if not wb.sheetnames:
            raise SpreadsheetParsingError("The Excel workbook contains no worksheets.")

        ws = wb.active
        if ws is None:
            raise SpreadsheetParsingError("The Excel workbook has no active worksheet.")

        rows: list[list[Any]] = []
        for row in ws.iter_rows(values_only=True):
            rows.append(list(row))
        return rows

    def _read_csv_rows(self, file_bytes: bytes) -> list[list[Any]]:
        """Reads rows from CSV bytes, handling various text encodings."""
        text = None
        for encoding in ("utf-8-sig", "utf-8", "latin-1"):
            try:
                text = file_bytes.decode(encoding)
                break
            except UnicodeDecodeError:
                continue

        if text is None:
            raise InvalidFileFormatError("Unable to decode CSV text. File must be encoded in UTF-8 or compatible text format.")

        reader = csv.reader(io.StringIO(text))
        return [row for row in reader]

    def _extract_and_validate_recipients(self, raw_rows: list[list[Any]]) -> list[RecipientInput]:
        """
        Locates headers, extracts row data, and performs rigorous validation.
        """
        # Find first non-empty row as header
        header_row_idx = -1
        header_cells: list[Any] = []
        for idx, row in enumerate(raw_rows):
            if any(cell is not None and str(cell).strip() != "" for cell in row):
                header_row_idx = idx
                header_cells = row
                break

        if header_row_idx == -1 or not header_cells:
            raise SpreadsheetParsingError("The uploaded spreadsheet is empty and contains no headers or records.")

        name_col_idx: int | None = None
        email_col_idx: int | None = None
        remarks_col_idx: int | None = None

        cleaned_headers = [_clean_header(c) for c in header_cells]

        for col_idx, cleaned in enumerate(cleaned_headers):
            if not cleaned:
                continue
            if name_col_idx is None and (cleaned in {_clean_header(v) for v in NAME_VARIANTS}):
                name_col_idx = col_idx
            elif email_col_idx is None and (cleaned in {_clean_header(v) for v in EMAIL_VARIANTS}):
                email_col_idx = col_idx
            elif remarks_col_idx is None and (cleaned in {_clean_header(v) for v in REMARKS_VARIANTS}):
                remarks_col_idx = col_idx

        if name_col_idx is None:
            found = [str(h) for h in header_cells if h is not None and str(h).strip()]
            raise SpreadsheetParsingError(
                f"Missing required 'recipient_name' (or 'name') column in header. Found columns: {found}"
            )

        recipients: list[RecipientInput] = []
        seen_recipients: set[tuple[str, str]] = set()

        # Parse subsequent rows
        data_rows = raw_rows[header_row_idx + 1 :]
        for offset, row in enumerate(data_rows, start=1):
            excel_row_num = header_row_idx + 1 + offset

            # Check if row is completely empty
            if not any(cell is not None and str(cell).strip() != "" for cell in row):
                continue

            # Extract Name
            raw_name = row[name_col_idx] if name_col_idx < len(row) else None
            if raw_name is None or not str(raw_name).strip():
                raise SpreadsheetParsingError(
                    f"Row {excel_row_num}: Recipient name is required and cannot be empty."
                )

            name = str(raw_name).strip()
            if len(name) < 2:
                raise SpreadsheetParsingError(
                    f"Row {excel_row_num}: Recipient name '{name}' must be at least 2 characters long."
                )
            if len(name) > 255:
                raise SpreadsheetParsingError(
                    f"Row {excel_row_num}: Recipient name exceeds maximum allowable length of 255 characters."
                )

            # Extract Email (optional)
            email: str | None = None
            if email_col_idx is not None and email_col_idx < len(row):
                raw_email = row[email_col_idx]
                if raw_email is not None:
                    cleaned_email = str(raw_email).strip()
                    if cleaned_email and cleaned_email.lower() not in ("none", "null", "n/a", "na", "-"):
                        email = cleaned_email

            # Extract Custom Remarks (optional)
            remarks: str | None = None
            if remarks_col_idx is not None and remarks_col_idx < len(row):
                raw_remarks = row[remarks_col_idx]
                if raw_remarks is not None:
                    cleaned_remarks = str(raw_remarks).strip()
                    if cleaned_remarks and cleaned_remarks.lower() not in ("none", "null", "n/a", "na", "-"):
                        remarks = cleaned_remarks[:500]

            # Validate through Pydantic RecipientInput schema
            try:
                recipient = RecipientInput(
                    recipient_name=name,
                    recipient_email=email,
                    custom_remarks=remarks,
                )
            except ValidationError as val_err:
                errors = val_err.errors()
                err_msg = errors[0].get("msg", "Invalid field format") if errors else "Validation failed"
                field_name = errors[0].get("loc", ["field"])[-1] if errors else "field"
                raise SpreadsheetParsingError(
                    f"Row {excel_row_num}: Validation failed for '{field_name}' - {err_msg}"
                ) from val_err

            # Validate batch uniqueness
            email_key = str(recipient.recipient_email).lower().strip() if recipient.recipient_email else ""
            key = (recipient.recipient_name.lower().strip(), email_key)
            if key in seen_recipients:
                raise SpreadsheetParsingError(
                    f"Row {excel_row_num}: Duplicate recipient detected: '{recipient.recipient_name}' "
                    f"({recipient.recipient_email or 'no email'}). Each recipient within a batch must be unique."
                )
            seen_recipients.add(key)
            recipients.append(recipient)

        if not recipients:
            raise SpreadsheetParsingError("The uploaded spreadsheet does not contain any recipient rows.")

        if len(recipients) > settings.MAX_BATCH_SIZE:
            raise BatchSizeExceededError(actual_size=len(recipients), max_size=settings.MAX_BATCH_SIZE)

        return recipients

    def generate_sample_recipients_excel(self) -> bytes:
        """
        Creates a pre-formatted, styled Excel (.xlsx) template with sample recipients.
        Useful for users downloading a template directly from Swagger UI or API.
        """
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Recipients"

        headers = ["recipient_name", "recipient_email", "custom_remarks"]
        ws.append(headers)

        sample_rows = [
            ("Ada Lovelace", "ada@example.com", "First Programmer"),
            ("Alan Turing", "alan@example.com", "Father of Computing"),
            ("Grace Hopper", "grace@example.com", "Compiler Pioneer"),
            ("Margaret Hamilton", "margaret@example.com", "Apollo Lead Software Engineer"),
            ("Dr. Alexander Montgomery Bartholomew-Fitzgerald III", "alex@example.com", "Distinguished Scholar"),
        ]

        for row in sample_rows:
            ws.append(list(row))

        # Styling
        header_fill = PatternFill(start_color="1A365D", end_color="1A365D", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        thin_border = Border(
            left=Side(style="thin", color="CBD5E0"),
            right=Side(style="thin", color="CBD5E0"),
            top=Side(style="thin", color="CBD5E0"),
            bottom=Side(style="thin", color="CBD5E0"),
        )

        for col_idx in range(1, 4):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="left", vertical="center")
            cell.border = thin_border

        # Set column widths
        ws.column_dimensions["A"].width = 45
        ws.column_dimensions["B"].width = 30
        ws.column_dimensions["C"].width = 35

        # Style data cells
        for row_idx in range(2, len(sample_rows) + 2):
            for col_idx in range(1, 4):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.border = thin_border
                cell.alignment = Alignment(vertical="center")

        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        return buffer.getvalue()


spreadsheet_parser = SpreadsheetParser()
